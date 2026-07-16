import os
from openpyxl import load_workbook
from django.db import transaction
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status

from employees.models import Employee, Department
from documents.serializers import EmployeeImportSerializer
from rest_framework.permissions import IsAuthenticated
from django.utils.decorators import method_decorator
from django.views.decorators.cache import cache_page

from api.services import EmployeeService
from api.serializers import EmployeeSerializer, EmployeeListSerializer
from api.responses import SuccessResponse, ErrorResponse
from api.permissions import IsHROrAdmin
from audit_logs.utils import create_audit_log, get_client_ip


class EmployeeListCreateView(APIView):
    """
    GET  /api/v1/employees/ — List all employees
    POST /api/v1/employees/ — Create a new employee
    """
    permission_classes = [IsAuthenticated]

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.service = EmployeeService()

    @method_decorator(cache_page(60 * 5))
    def get(self, request):
        employees = self.service.list_employees()
        serializer = EmployeeListSerializer(employees, many=True)
        return SuccessResponse(
            data=serializer.data,
            message="Employees retrieved successfully",
        )

    def post(self, request):
        serializer = EmployeeSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        employee = serializer.save()

        # Audit log with user context
        create_audit_log(
            user=request.user,
            action="CREATE",
            module="EMPLOYEE",
            object_id=employee.pk,
            details=serializer.data,
            ip_address=get_client_ip(request),
        )

        return SuccessResponse(
            data=EmployeeSerializer(employee).data,
            message="Employee created successfully",
            status_code=status.HTTP_201_CREATED,
        )


class EmployeeDetailView(APIView):
    """
    GET    /api/v1/employees/{id}/ — Retrieve an employee
    PUT    /api/v1/employees/{id}/ — Update an employee
    DELETE /api/v1/employees/{id}/ — Delete an employee
    """
    permission_classes = [IsAuthenticated]

    def __init__(self, **kwargs):
        super().__init__(**kwargs)
        self.service = EmployeeService()

    def get(self, request, pk):
        employee = self.service.get_employee(pk)
        serializer = EmployeeSerializer(employee)
        return SuccessResponse(
            data=serializer.data,
            message="Employee retrieved successfully",
        )

    def put(self, request, pk):
        employee = self.service.get_employee(pk)
        serializer = EmployeeSerializer(employee, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        updated_employee = serializer.save()

        # Audit log
        create_audit_log(
            user=request.user,
            action="UPDATE",
            module="EMPLOYEE",
            object_id=updated_employee.pk,
            details=serializer.validated_data,
            ip_address=get_client_ip(request),
        )

        return SuccessResponse(
            data=EmployeeSerializer(updated_employee).data,
            message="Employee updated successfully",
        )

    def delete(self, request, pk):
        employee = self.service.get_employee(pk)
        employee_data = EmployeeSerializer(employee).data
        self.service.delete_employee(pk)

        # Audit log
        create_audit_log(
            user=request.user,
            action="DELETE",
            module="EMPLOYEE",
            object_id=pk,
            details=employee_data,
            ip_address=get_client_ip(request),
        )

        return SuccessResponse(
            message="Employee deleted successfully",
            status_code=status.HTTP_200_OK,
        )
class EmployeeImportView(APIView):
    """
    Import employees from Excel (.xlsx)
    """
    permission_classes = [IsHROrAdmin]

    def post(self, request):
        serializer = EmployeeImportSerializer(data=request.data)

        if not serializer.is_valid():
            return ErrorResponse(message="Validation Error", errors=serializer.errors, status_code=status.HTTP_400_BAD_REQUEST)

        excel_file = request.FILES.get('file')
        try:
            wb = load_workbook(excel_file, data_only=True)
            sheet = wb.active
        except Exception as e:
            return ErrorResponse(message="Invalid Excel file", errors=str(e), status_code=status.HTTP_400_BAD_REQUEST)

        total_rows = 0
        created = 0
        failed = 0
        errors = []

        for idx, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):  # skip empty rows
                continue
            
            total_rows += 1

            try:
                emp_id, first_name, last_name, email, phone, dept_name, designation, salary, joining_date = row[:9]
                
                if not all([emp_id, first_name, email, dept_name]):
                    failed += 1
                    errors.append({"row": idx, "error": "Missing required fields (ID, Name, Email, Dept)"})
                    continue

                department, _ = Department.objects.get_or_create(name=dept_name)

                with transaction.atomic():
                    Employee.objects.update_or_create(
                        employee_id=str(emp_id),
                        defaults={
                            "first_name": str(first_name),
                            "last_name": str(last_name) if last_name else "",
                            "email": str(email),
                            "phone": str(phone) if phone else "",
                            "department": department,
                            "designation": str(designation) if designation else "",
                            "salary": float(salary) if salary else 0.0,
                            "joining_date": joining_date,
                        }
                    )
                created += 1

            except Exception as e:
                failed += 1
                errors.append({"row": idx, "error": str(e)})

        # Audit log
        create_audit_log(
            user=request.user,
            action="CREATE",
            module="EMPLOYEE",
            object_id=0,
            details={"imported": created, "failed": failed},
            ip_address=get_client_ip(request),
        )

        return Response(
            {
                "total_rows": total_rows,
                "created": created,
                "failed": failed,
                "errors": errors
            },
            status=status.HTTP_200_OK,
        )

from reports.excel_export import generate_employees_excel

class EmployeeExportView(APIView):
    """
    GET /api/v1/employees/export/ — Export employee data to Excel
    """
    permission_classes = [IsHROrAdmin]

    def get(self, request):
        export_type = request.GET.get('type', 'employee')
        
        try:
            workbook_path = generate_employees_excel(export_type)
            from django.http import FileResponse
            response = FileResponse(open(workbook_path, 'rb'), as_attachment=True, filename=f'{export_type}_export.xlsx')
            return response
        except Exception as e:
            return ErrorResponse(message="Export failed", errors=str(e), status_code=status.HTTP_500_INTERNAL_SERVER_ERROR)

from reports.csv_handler import import_employees_csv, export_employees_csv

class EmployeeCSVImportView(APIView):
    """
    POST /api/v1/employees/csv/import/ — Import employee data from CSV
    """
    permission_classes = [IsHROrAdmin]

    def post(self, request):
        csv_file = request.FILES.get('file')
        if not csv_file or not csv_file.name.endswith('.csv'):
            return ErrorResponse(message="Invalid file format. Please upload a .csv file.", status_code=status.HTTP_400_BAD_REQUEST)
        
        try:
            result = import_employees_csv(csv_file)
            
            # Audit log
            create_audit_log(
                user=request.user,
                action="CREATE",
                module="EMPLOYEE",
                object_id=0,
                details={"imported": result['created'], "failed": result['failed'], "type": "CSV"},
                ip_address=get_client_ip(request),
            )
            return Response(result, status=status.HTTP_200_OK)
        except Exception as e:
            return ErrorResponse(message="CSV Import failed", errors=str(e), status_code=status.HTTP_400_BAD_REQUEST)


class EmployeeCSVExportView(APIView):
    """
    GET /api/v1/employees/csv/export/ — Export employee data to CSV
    """
    permission_classes = [IsHROrAdmin]

    def get(self, request):
        try:
            csv_path = export_employees_csv()
            from django.http import FileResponse
            response = FileResponse(open(csv_path, 'rb'), as_attachment=True, filename='employees_export.csv')
            return response
        except Exception as e:
            return ErrorResponse(message="CSV Export failed", errors=str(e), status_code=status.HTTP_500_INTERNAL_SERVER_ERROR)

from reports.pdf_generator import generate_employee_profile_pdf, generate_salary_slip_pdf, generate_id_card_pdf

class EmployeeProfilePDFView(APIView):
    """
    GET /api/v1/employees/{id}/pdf-profile/
    """
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            employee = Employee.objects.get(pk=pk)
            pdf_path = generate_employee_profile_pdf(employee)
            from django.http import FileResponse
            return FileResponse(open(pdf_path, 'rb'), as_attachment=True, filename=os.path.basename(pdf_path))
        except Employee.DoesNotExist:
            return ErrorResponse(message="Employee not found", status_code=status.HTTP_404_NOT_FOUND)
        except Exception as e:
            return ErrorResponse(message="PDF generation failed", errors=str(e), status_code=status.HTTP_500_INTERNAL_SERVER_ERROR)

class EmployeeSalarySlipPDFView(APIView):
    """
    GET /api/v1/employees/{id}/salary-slip/?month=June&year=2026
    """
    permission_classes = [IsHROrAdmin]

    def get(self, request, pk):
        month = request.GET.get('month', 'Current')
        year = request.GET.get('year', 'Year')
        try:
            employee = Employee.objects.get(pk=pk)
            pdf_path = generate_salary_slip_pdf(employee, month, year)
            from django.http import FileResponse
            return FileResponse(open(pdf_path, 'rb'), as_attachment=True, filename=os.path.basename(pdf_path))
        except Employee.DoesNotExist:
            return ErrorResponse(message="Employee not found", status_code=status.HTTP_404_NOT_FOUND)
        except Exception as e:
            return ErrorResponse(message="Salary slip generation failed", errors=str(e), status_code=status.HTTP_500_INTERNAL_SERVER_ERROR)

class EmployeeIDCardPDFView(APIView):
    """
    GET /api/v1/employees/{id}/id-card/
    """
    permission_classes = [IsHROrAdmin]

    def get(self, request, pk):
        try:
            employee = Employee.objects.get(pk=pk)
            pdf_path = generate_id_card_pdf(employee)
            from django.http import FileResponse
            return FileResponse(open(pdf_path, 'rb'), as_attachment=True, filename=os.path.basename(pdf_path))
        except Employee.DoesNotExist:
            return ErrorResponse(message="Employee not found", status_code=status.HTTP_404_NOT_FOUND)
        except Exception as e:
            return ErrorResponse(message="ID card generation failed", errors=str(e), status_code=status.HTTP_500_INTERNAL_SERVER_ERROR)

